"""PPTX 文本抽取（可选依赖 python-pptx）。"""

from __future__ import annotations

from pathlib import Path

MAX_SLIDES = 15


def extract_pptx_text(path: str | Path, *, max_chars: int = 2000) -> str:
    """使用 python-pptx 抽取幻灯片文本。未安装时抛 ImportError。"""
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ImportError("PPTX 抽取需要 python-pptx：pip install python-pptx") from exc

    prs = Presentation(str(path))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides):
        if i >= MAX_SLIDES:
            break
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                parts.append(text.strip())
    return "\n".join(parts)[:max_chars]
